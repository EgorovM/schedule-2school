# -*- coding: utf-8 -*-
"""Импорт рабочего файла завуча (xlsx) в модель расписания.

Читает:
  «РАСП С БАЛЛ»          — расписание 5—11 классов с баллами трудности (основной источник);
  «расписание нач классы» — 1—4 классы;
  «список учителей»       — справочник педработников;
  «прием»                 — прошлогодняя расстановка учителей (для подсказок).
"""
from __future__ import annotations

import re
from pathlib import Path

from . import sanpin

ROOM_PREFIX = re.compile(r"^((?:\d{1,3})(?:\s*[/(\[|]\s*\d{1,3}\s*[)\]]?)*)\s+(?=\D)")
CANON = {
    "русский язык": "Русский язык", "литературное чтение": "Литературное чтение",
    "литература": "Литература", "математика": "Математика",
    "окружающий мир": "Окружающий мир", "окр мир": "Окружающий мир",
    "окр мир 1": "Окружающий мир", "изо": "ИЗО", "музыка": "Музыка",
    "труд": "Труд (технология)", "технология": "Труд (технология)",
    "труд (технология)": "Труд (технология)",
    "физическая культура": "Физическая культура", "физкультура": "Физическая культура",
    "иностранный язык": "Иностранный язык", "англ яз": "Иностранный язык",
    "английский язык": "Иностранный язык", "орксэ": "ОРКСЭ",
    "вуд яягос": "ВУД «Якутский язык»", "яягос": "ЯЯГОС",
    "обзр": "ОБЗР", "алгебра": "Алгебра", "геометрия": "Геометрия",
    "информатика": "Информатика", "биология": "Биология", "физика": "Физика",
    "химия": "Химия", "география": "География", "история": "История",
    "обществознание": "Обществознание", "родной язык": "Родной язык",
    "родная литература": "Родная литература",
    "вероятность и статистика": "Вероятность и статистика",
}


def txt(v) -> str:
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip() if v is not None else ""


def canon(name: str) -> str:
    key = name.lower().strip(" .")
    if key in CANON:
        return CANON[key]
    return name[:1].upper() + name[1:] if name else name


def split_room(s: str) -> tuple[str, str]:
    """«46 Физическая культура» → ('46', 'Физическая культура')."""
    m = ROOM_PREFIX.match(s)
    if m:
        return re.sub(r"\s+", "", m.group(1)), s[m.end():].strip()
    return "", s


def parse_teachers(cell: str) -> list[dict]:
    """Разбирает ячейку листа «прием»: номера учителей + предмет.

    «45(49) Русский язык»              → [{'subject': 'Русский язык', 'ids': ['45', '49']}]
    «20 иностранный язык/1 Информатика» → две записи (класс делится на подгруппы)
    «22/30 Русский язык»                → один предмет, два учителя
    """
    s = txt(cell)
    if not s:
        return []
    out: list[dict] = []
    pending: list[str] = []
    for part in re.split(r"[/|]", s):
        part = part.strip()
        if not part:
            continue
        ids = re.findall(r"\d+", re.match(r"^[\d\s()\[\]]*", part).group() or "")
        rest = re.sub(r"^[\d\s()\[\]]*", "", part).strip()
        if rest:
            out.append({"subject": canon(rest), "ids": pending + ids})
            pending = []
        else:
            pending += ids
    if pending and out:
        out[0]["ids"] = pending + out[0]["ids"]
    elif pending:
        out.append({"subject": "", "ids": pending})
    return out


def teacher_code(fio: str) -> str:
    parts = [p for p in re.split(r"\s+", fio) if p]
    return "".join(p[0].upper() for p in parts[:3])


def _class_grade(cls: str) -> int:
    m = re.match(r"\d+", cls)
    return int(m.group()) if m else 0


def _read_block(ws, day_col, num_col, time_col, first_col, last_col, step, shift):
    """Читает один блок расписания (смену). step=2 — пары «предмет, балл»."""
    classes: list[dict] = []
    for c in range(first_col, last_col + 1, step):
        name = txt(ws.cell(row=8, column=c).value).replace(" ", "")
        if not name or name in ("Дни", "№", "Время"):
            continue
        classes.append({"id": name, "col": c, "grade": _class_grade(name),
                        "shift": shift, "room": txt(ws.cell(row=9, column=c).value)})

    lessons, bells, day = [], {}, None
    for row in range(10, ws.max_row + 1):
        d = txt(ws.cell(row=row, column=day_col).value)
        if d:
            day = sanpin.FULL_DAY.get(d.lower())
        n = txt(ws.cell(row=row, column=num_col).value)
        if not day or not n.isdigit():
            continue
        n = int(n)
        time = txt(ws.cell(row=row, column=time_col).value).replace(".", ":")
        time = re.sub(r"\s*-\s*", "—", time)
        if time and n not in bells:
            bells[n] = time
        for c in classes:
            raw = txt(ws.cell(row=row, column=c["col"]).value)
            if not raw or raw.isdigit():
                continue
            room, subject = split_room(raw)
            score = None
            if step == 2:
                sv = txt(ws.cell(row=row, column=c["col"] + 1).value)
                score = int(sv) if sv.isdigit() else None
            lessons.append({"cls": c["id"], "day": day, "n": n,
                            "subject": canon(subject), "room": room or None,
                            "score_file": score, "teacher": None})
    for c in classes:
        c.pop("col", None)
    return classes, lessons, bells


def import_workbook(path: str | Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)

    ws = wb["РАСП С БАЛЛ"]
    title = txt(ws.cell(row=5, column=1).value)
    director = txt(ws.cell(row=3, column=3).value).strip("_ /")

    cls1, les1, bells1 = _read_block(ws, 1, 2, 3, 4, 21, 2, shift=1)
    cls2, les2, bells2 = _read_block(ws, 22, 23, 24, 25, 42, 2, shift=2)

    wp = wb["расписание нач классы"]
    pclasses, plessons = [], []
    for c in range(5, 13):
        name = txt(wp.cell(row=4, column=c).value).replace(" ", "")
        if name:
            pclasses.append({"id": name, "col": c, "grade": _class_grade(name),
                             "shift": 1, "room": ""})
    day = None
    for row in range(5, wp.max_row + 1):
        d = txt(wp.cell(row=row, column=3).value)
        if d:
            day = sanpin.FULL_DAY.get(d.lower())
        n = txt(wp.cell(row=row, column=4).value)
        if not day or not n.isdigit():
            continue
        for c in pclasses:
            raw = txt(wp.cell(row=row, column=c["col"]).value)
            if not raw:
                continue
            for item in parse_teachers(raw) or [{"subject": canon(raw), "ids": []}]:
                if not item["subject"]:
                    continue
                plessons.append({"cls": c["id"], "day": day, "n": int(n),
                                 "subject": item["subject"], "room": None,
                                 "score_file": None,
                                 "teacher": item["ids"][0] if len(item["ids"]) == 1 else None,
                                 "teachers": item["ids"] or None})
    for c in pclasses:
        c.pop("col", None)

    teachers = []
    wt = wb["список учителей "] if "список учителей " in wb.sheetnames else wb["список учителей"]
    for r in range(4, wt.max_row + 1):
        num = txt(wt.cell(row=r, column=1).value)
        fio = txt(wt.cell(row=r, column=2).value)
        subj = txt(wt.cell(row=r, column=3).value)
        if not num.isdigit() or len(fio.split()) < 3:
            continue           # хвост листа содержит не педработников
        teachers.append({"id": num, "name": fio, "code": teacher_code(fio),
                         "subjects": [s.strip() for s in re.split(r"[,;]", subj) if s.strip()]})

    prior = []
    if "прием" in wb.sheetnames:
        wa = wb["прием"]
        cols = {}
        for c in range(4, wa.max_column + 1):
            name = txt(wa.cell(row=8, column=c).value).replace(" ", "")
            if name and name not in ("Дни", "№", "Время"):
                cols[c] = name
        day = None
        for row in range(10, wa.max_row + 1):
            for dc in (1, 13, 22):
                d = txt(wa.cell(row=row, column=dc).value)
                if d and d.lower() in sanpin.FULL_DAY:
                    day = sanpin.FULL_DAY[d.lower()]
            n = txt(wa.cell(row=row, column=2).value)
            if not day or not n.isdigit():
                continue
            for c, cls in cols.items():
                for item in parse_teachers(wa.cell(row=row, column=c).value):
                    if item["subject"]:
                        prior.append({"cls": cls, "day": day, "n": int(n),
                                      "subject": item["subject"], "ids": item["ids"]})

    classes = pclasses + cls1 + cls2
    for c in classes:
        c["days"] = 5 if c["grade"] <= 4 else 6
    all_lessons = plessons + les1 + les2
    for i, l in enumerate(all_lessons, 1):
        l["id"] = i
    return {
        "meta": {"title": title, "director": director, "year": "2026/2027",
                 "source": str(path), "week_days": 6},
        "classes": classes,
        "bells": {"1": bells1, "2": bells2},
        "teachers": teachers,
        "lessons": all_lessons,
        "prior": prior,
    }
