# -*- coding: utf-8 -*-
"""Локальный сервер инструмента. Запуск: python3 -m raspisanie"""
from __future__ import annotations

import json
import re
import webbrowser
from datetime import date, datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from . import plan as planner
from . import rules, sanpin, solver, store

STATIC = Path(__file__).resolve().parent / "static"
WEEKDAY = {0: "пн", 1: "вт", 2: "ср", 3: "чт", 4: "пт", 5: "сб"}


def state() -> dict:
    model = store.load_model()
    scores = store.load_scores()
    analysis = rules.analyze(model, scores)
    # балл трудности — величина вычисляемая, в файл не пишется
    for l in model["lessons"]:
        grade = int(re.match(r"\d+", l["cls"]).group()) if re.match(r"\d+", l["cls"]) else 0
        l["score"] = None if rules.is_extracurricular(l["subject"]) \
            else rules.score_of(l, grade, scores)
    return {"model": model, "analysis": analysis, "scores": scores,
            "subs": store.load_subs(), "prefs": store.load_prefs(),
            "plan": planner.build(model), "load": planner.teacher_load(model)}


def lesson_key(l: dict) -> tuple:
    return (l["cls"], l["day"], l["n"])


def day_sheet(model: dict, subs: dict, on: str) -> dict:
    """Лист на конкретный день с учётом замен."""
    d = datetime.strptime(on, "%Y-%m-%d").date()
    day = WEEKDAY.get(d.weekday())
    absent = {a["teacher"] for a in subs["absences"] if a["date"] == on}
    repl = {(s["cls"], s["day"], s["n"]): s for s in subs["substitutions"] if s["date"] == on}
    teachers = {t["id"]: t for t in model["teachers"]}
    rows = []
    for l in model["lessons"]:
        if l["day"] != day:
            continue
        s = repl.get(lesson_key(l))
        rows.append({
            **l,
            "teacher_name": teachers.get(l.get("teacher"), {}).get("name"),
            "absent": bool(l.get("teacher") and l["teacher"] in absent),
            "replacement": s,
            "replacement_name": teachers.get((s or {}).get("to"), {}).get("name"),
        })
    rows.sort(key=lambda r: (r["cls"], r["n"]))
    return {"date": on, "day": day, "day_full": sanpin.DAY_FULL.get(day, ""),
            "rows": rows, "absent": sorted(absent)}


PREVIEW: dict = {}


class Handler(BaseHTTPRequestHandler):
    server_version = "raspisanie/1.0"

    def log_message(self, fmt, *args):
        pass

    # ---------- helpers ----------
    def _send(self, payload, code=200, ctype="application/json; charset=utf-8"):
        body = payload if isinstance(payload, bytes) else \
            json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def _body(self) -> dict:
        n = int(self.headers.get("Content-Length") or 0)
        return json.loads(self.rfile.read(n) or b"{}")

    # ---------- GET ----------
    def do_GET(self):
        u = urlparse(self.path)
        q = parse_qs(u.query)
        path = u.path

        if path in ("/", "/index.html"):
            return self._file("index.html", "text/html; charset=utf-8")
        if path.startswith("/static/"):
            name = path[len("/static/"):]
            ctype = {"js": "application/javascript; charset=utf-8",
                     "css": "text/css; charset=utf-8"}.get(name.rsplit(".", 1)[-1], "text/plain")
            return self._file(name, ctype)

        if path == "/api/state":
            return self._send(state())

        if path == "/api/free":
            model = store.load_model()
            free = rules.free_teachers(model, q.get("day", [""])[0],
                                       int(q.get("n", ["0"])[0]),
                                       q.get("subject", [None])[0])
            return self._send({"teachers": free})

        if path == "/api/suggest":
            model = store.load_model()
            return self._send({"teachers": rules.suggest_teacher(
                model, q.get("cls", [""])[0], q.get("subject", [""])[0])})

        if path == "/api/day":
            on = q.get("date", [date.today().isoformat()])[0]
            return self._send(day_sheet(store.load_model(), store.load_subs(), on))

        if path == "/api/preview":
            return self._send(PREVIEW or {"empty": True})

        if path == "/api/export":
            return self._export(q.get("kind", ["xlsx"])[0])

        return self._send({"error": "не найдено"}, 404)

    def _file(self, name: str, ctype: str):
        p = STATIC / name
        if not p.exists():
            return self._send({"error": "нет файла"}, 404)
        return self._send(p.read_bytes(), ctype=ctype)

    # ---------- POST ----------
    def do_POST(self):
        path = urlparse(self.path).path
        try:
            body = self._body()
        except Exception as e:
            return self._send({"error": f"не разобрал запрос: {e}"}, 400)

        if path == "/api/import":
            src = body.get("path")
            if not src or not Path(src).exists():
                return self._send({"error": "файл не найден"}, 400)
            from .importer import import_workbook
            model = import_workbook(src)
            store.save_model(model, f"импорт из {Path(src).name}")
            return self._send(state())

        model = store.load_model()

        if path == "/api/move":
            return self._move(model, body)
        if path == "/api/assign":
            lid = body.get("id")
            for l in model["lessons"]:
                if l["id"] == lid:
                    l["teacher"] = body.get("teacher") or None
                    store.save_model(model, f"учитель для {l['cls']} {l['day']}/{l['n']}")
                    return self._send(state())
            return self._send({"error": "урок не найден"}, 404)

        if path == "/api/assign-bulk":
            cls, subject, teacher = body.get("cls"), body.get("subject"), body.get("teacher")
            k = 0
            for l in model["lessons"]:
                if l["cls"] == cls and l["subject"] == subject:
                    l["teacher"] = teacher or None
                    k += 1
            store.save_model(model, f"{teacher or '—'} → {cls} «{subject}» ({k})")
            return self._send(state())

        if path == "/api/plan/autoassign":
            res = planner.autoassign(model)
            store.save_model(model, "автоназначение учителей")
            return self._send({**state(), "autoassign": res})

        if path == "/api/plan/teacher":
            cls, subject, tid = body["cls"], body["subject"], body.get("teacher") or None
            for l in model["lessons"]:
                if l["cls"] == cls and l["subject"] == subject:
                    l["teacher"] = tid
            store.save_model(model, f"учитель для {cls} «{subject}»")
            return self._send(state())

        if path == "/api/plan/hours":
            cls, subject, hours = body["cls"], body["subject"], int(body["hours"])
            mine = [l for l in model["lessons"] if l["cls"] == cls and l["subject"] == subject]
            if hours < len(mine):
                for l in mine[hours:]:
                    model["lessons"].remove(l)
            elif hours > len(mine) and mine:
                nxt = max((l["id"] for l in model["lessons"]), default=0) + 1
                for k in range(hours - len(mine)):
                    model["lessons"].append({**mine[0], "id": nxt + k, "day": mine[0]["day"],
                                             "n": mine[0]["n"]})
            store.save_model(model, f"{cls} «{subject}»: {hours} ч")
            return self._send(state())

        if path == "/api/prefs":
            store.save_prefs(body.get("prefs") or [])
            return self._send(state())

        if path == "/api/generate":
            global PREVIEW
            seconds = max(3.0, min(120.0, float(body.get("seconds", 20))))
            attempts = max(1, min(5, int(body.get("attempts", 1))))
            scores = store.load_scores()
            lessons, report = solver.build(model, scores, store.load_prefs(),
                                           {"seconds": seconds, "attempts": attempts,
                                            "seed": int(body.get("seed", 1))})
            preview_model = {**model, "lessons": lessons}
            PREVIEW = {"lessons": lessons, "report": report,
                       "analysis": rules.analyze(preview_model, scores),
                       "before": rules.analyze(model, scores)["summary"]}
            return self._send(PREVIEW)

        if path == "/api/apply":
            if not PREVIEW:
                return self._send({"error": "сначала соберите расписание"}, 400)
            model["lessons"] = PREVIEW["lessons"]
            store.save_model(model, "расписание собрано конструктором")
            PREVIEW = {}
            return self._send(state())

        if path == "/api/scores":
            store.save_scores(body.get("scores") or {})
            return self._send(state())

        if path == "/api/absence":
            subs = store.load_subs()
            rec = {"date": body["date"], "teacher": body["teacher"],
                   "reason": body.get("reason", "")}
            subs["absences"] = [a for a in subs["absences"]
                                if not (a["date"] == rec["date"] and a["teacher"] == rec["teacher"])]
            if not body.get("remove"):
                subs["absences"].append(rec)
            store.save_subs(subs)
            return self._send({"ok": True, "subs": subs})

        if path == "/api/substitute":
            subs = store.load_subs()
            key = (body["cls"], body["day"], body["n"], body["date"])
            subs["substitutions"] = [
                s for s in subs["substitutions"]
                if (s["cls"], s["day"], s["n"], s["date"]) != key]
            if not body.get("remove"):
                subs["substitutions"].append({
                    "date": body["date"], "cls": body["cls"], "day": body["day"],
                    "n": body["n"], "subject": body.get("subject"),
                    "from": body.get("from"), "to": body.get("to"),
                    "kind": body.get("kind", "замена"),
                    "created": datetime.now().isoformat(timespec="minutes")})
            store.save_subs(subs)
            return self._send({"ok": True, "subs": subs})

        return self._send({"error": "не найдено"}, 404)

    def _move(self, model, body):
        """Перенос урока в пустой слот или обмен двух уроков местами."""
        lid, day, n = body.get("id"), body.get("day"), int(body.get("n"))
        src = next((l for l in model["lessons"] if l["id"] == lid), None)
        if not src:
            return self._send({"error": "урок не найден"}, 404)
        dst = [l for l in model["lessons"]
               if l["cls"] == src["cls"] and l["day"] == day and l["n"] == n]
        if dst and len(dst) == 1:
            other = dst[0]
            other["day"], other["n"], src["day"], src["n"] = src["day"], src["n"], day, n
            reason = f"{src['cls']}: «{src['subject']}» ↔ «{other['subject']}»"
        elif dst:
            return self._send({"error": "в этом слоте несколько уроков — разведите подгруппы"}, 409)
        else:
            src["day"], src["n"] = day, n
            reason = f"{src['cls']}: «{src['subject']}» → {day}/{n}"
        store.save_model(model, reason)
        return self._send(state())

    # ---------- экспорт ----------
    def _export(self, kind: str):
        model = store.load_model()
        if kind == "json":
            return self._send(model)
        try:
            import openpyxl
        except ImportError:
            return self._send({"error": "нет openpyxl"}, 500)
        from openpyxl.styles import Alignment, Border, Font, Side
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Расписание"
        thin = Side(style="thin", color="B0B0B0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        teachers = {t["id"]: t for t in model["teachers"]}
        classes = [c for c in model["classes"]]
        ws.cell(row=1, column=1, value=model["meta"].get("title", "Расписание уроков")).font = Font(bold=True)
        ws.cell(row=2, column=1, value=f"Выгружено {datetime.now():%d.%m.%Y %H:%M}")
        head = ["Класс", "День", "Урок", "Предмет", "Учитель", "Кабинет"]
        for j, h in enumerate(head, 1):
            c = ws.cell(row=4, column=j, value=h)
            c.font = Font(bold=True)
            c.border = border
        order = {d: i for i, d in enumerate(sanpin.DAYS)}
        rows = sorted(model["lessons"],
                      key=lambda l: ([c["id"] for c in classes].index(l["cls"])
                                     if any(c["id"] == l["cls"] for c in classes) else 99,
                                     order.get(l["day"], 9), l["n"]))
        r = 5
        for l in rows:
            vals = [l["cls"], sanpin.DAY_FULL.get(l["day"], l["day"]), l["n"], l["subject"],
                    teachers.get(l.get("teacher"), {}).get("name", ""), l.get("room") or ""]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=j, value=v)
                c.border = border
                c.alignment = Alignment(vertical="center")
            r += 1
        for col, w in zip("ABCDEF", (8, 14, 7, 28, 32, 10)):
            ws.column_dimensions[col].width = w
        import io
        buf = io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()
        self.send_response(200)
        self.send_header("Content-Type",
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition",
                         'attachment; filename="raspisanie.xlsx"')
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)


def run(port: int = 8770, open_browser: bool = True) -> None:
    srv = ThreadingHTTPServer(("127.0.0.1", port), Handler)
    url = f"http://127.0.0.1:{port}/"
    print(f"Инструмент расписания: {url}\nОстановить — Ctrl+C")
    if open_browser:
        webbrowser.open(url)
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("\nостановлен")
